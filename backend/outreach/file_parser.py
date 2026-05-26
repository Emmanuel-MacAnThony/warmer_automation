"""
Parse uploaded template files (TXT, DOCX, PDF) and detect [variable] placeholders.

Returns raw text and the list of unique placeholder names found — no LLM involved.
"""
from __future__ import annotations

import io
import re
from typing import Any


# Matches [word] and [word/word] style placeholders (not [word | fallback:...] — those are resolved slots)
_BRACKET_RE = re.compile(r'\[([A-Za-z][A-Za-z0-9 _/.-]*)\]')


def _extract_text_txt(data: bytes) -> str:
    return data.decode('utf-8', errors='replace')


def _extract_text_docx(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_text_pdf(data: bytes) -> str:
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return '\n\n'.join(pages)


def parse_file(filename: str, data: bytes) -> dict[str, Any]:
    """
    Extract text and detect bracket placeholders from an uploaded file.

    Returns
    -------
    {
      'text': str,                  — raw extracted text
      'detected_vars': list[str],   — unique placeholder names found, e.g. ['Name', 'sector/geography']
    }
    """
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''

    if ext == 'pdf':
        text = _extract_text_pdf(data)
    elif ext in ('docx', 'doc'):
        text = _extract_text_docx(data)
    else:
        text = _extract_text_txt(data)

    detected = list(dict.fromkeys(m.group(1) for m in _BRACKET_RE.finditer(text)))

    return {'text': text, 'detected_vars': detected}
